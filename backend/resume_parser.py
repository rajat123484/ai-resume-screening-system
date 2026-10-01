import os
import pymupdf
from docx import Document
from paddleocr import PaddleOCR


# =========================================================
# OCR MODEL
# =========================================================

ocr = PaddleOCR(
    lang="en"
)


# =========================================================
# MAIN RESUME TEXT EXTRACTOR
# =========================================================

def extract_text(file_path):

    extension = os.path.splitext(file_path)[1].lower()

    if extension == ".pdf":
        return extract_text_from_pdf(file_path)

    elif extension == ".docx":
        return extract_text_from_docx(file_path)

    else:
        return ""


# =========================================================
# PDF TEXT EXTRACTION
# =========================================================

def extract_text_from_pdf(file_path):

    doc = pymupdf.open(file_path)

    text = ""

    for page in doc:
        text += page.get_text()

    doc.close()

    # -----------------------------------------------------
    # If normal PDF extraction worked
    # -----------------------------------------------------

    if len(text.strip()) >= 50:
        return text

    # -----------------------------------------------------
    # Otherwise use OCR
    # -----------------------------------------------------

    return extract_text_with_ocr(file_path)


# =========================================================
# PDF OCR
# =========================================================

def extract_text_with_ocr(file_path):

    doc = pymupdf.open(file_path)

    extracted_text = ""

    for page_number, page in enumerate(doc):

        pixmap = page.get_pixmap(
            matrix=pymupdf.Matrix(2, 2)
        )

        image_path = f"temp_page_{page_number}.png"

        pixmap.save(image_path)

        try:

            result = ocr.predict(image_path)

            for res in result:

                if hasattr(res, "json"):

                    data = res.json

                    if callable(data):
                        data = data()

                    if isinstance(data, dict):

                        ocr_data = data.get(
                            "res",
                            data
                        )

                        texts = ocr_data.get(
                            "rec_texts",
                            []
                        )

                        extracted_text += "\n".join(
                            texts
                        )

            extracted_text += "\n"

        finally:

            if os.path.exists(image_path):
                os.remove(image_path)

    doc.close()

    return extracted_text


# =========================================================
# DOCX TEXT EXTRACTION
# =========================================================

def extract_text_from_docx(file_path):

    document = Document(file_path)

    extracted_text = []

    # -----------------------------------------------------
    # Paragraphs
    # -----------------------------------------------------

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            extracted_text.append(text)


    # -----------------------------------------------------
    # Tables
    # -----------------------------------------------------

    for table in document.tables:

        for row in table.rows:

            row_text = []

            for cell in row.cells:

                cell_text = cell.text.strip()

                if cell_text:
                    row_text.append(cell_text)

            if row_text:
                extracted_text.append(
                    " | ".join(row_text)
                )


    return "\n".join(extracted_text)